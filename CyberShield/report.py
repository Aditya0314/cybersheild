
from pathlib import Path
from datetime import datetime
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

BASE_DIR = Path(__file__).resolve().parent
LOG_FILE = BASE_DIR / "sample_log.txt"


def export_report():
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Security Report"

    # Report title
    sheet["A1"] = "CyberShield Security Report"
    sheet["A1"].font = Font(size=16, bold=True, color="FFFFFF")
    sheet["A1"].fill = PatternFill(
        "solid", fgColor="17365D"
    )

    sheet.merge_cells("A1:D1")

    # Column headings
    headers = ["Timestamp", "Username", "IP Address", "Status"]
    sheet.append(headers)

    for cell in sheet[2]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = PatternFill("solid", fgColor="1677B8")
        cell.alignment = Alignment(horizontal="center")

    # Read login records
    if LOG_FILE.exists():
        with open(LOG_FILE, "r", encoding="utf-8") as file:
            for line in file:
                parts = line.strip().split(" | ")

                if len(parts) != 4:
                    continue

                timestamp = parts[0]
                username = parts[1].replace("user=", "").strip()
                ip_address = parts[2].removeprefix("ip=")
                status = parts[3].removeprefix("status=")

                sheet.append([
                    timestamp, username, ip_address, status
                ])

    # Format columns
    for column, width in {
        "A": 16, "B": 18, "C": 20, "D": 16
    }.items():
        sheet.column_dimensions[column].width = width

    sheet.freeze_panes = "A3"
    sheet.auto_filter.ref = f"A2:D{sheet.max_row}"

    # Save the report
    report_path = BASE_DIR / "CyberShield_Report.xlsx"
    workbook.save(report_path)

    return report_path
